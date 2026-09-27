"""Parse a spreadsheet of phrases for bulk import.

Accepts .xlsx (openpyxl), legacy .xls (xlrd) and .csv. The expected shape is
one phrase per row:

    A: phrase text            (required)
    B: locale                 (optional, defaults to en-JM)
    C: active                 (optional, defaults to true)

A header row is detected and skipped, so a file exported from Excel with
"Phrase"/"Locale"/"Active" column titles works without editing.
"""
from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass, field
from typing import Optional

from ..models.phrase import PHRASE_TEXT_MAX_CHARS

logger = logging.getLogger("carib.phrase_import")

MIN_LEN = 3
# Follows the column, so a phrase is only ever rejected here for being
# unstorable — being too long to *read* is caught by the duration bands.
MAX_LEN = PHRASE_TEXT_MAX_CHARS
DEFAULT_LOCALE = "en-JM"

# First-cell values that mean "this row is a header, not a phrase".
_HEADER_CELLS = {"phrase", "phrases", "text", "phrase text", "sentence", "prompt"}
_FALSEY = {"0", "false", "no", "n", "inactive", "off"}


class PhraseImportError(ValueError):
    """The file could not be read at all."""


@dataclass
class ParsedPhrase:
    row: int
    text: str
    locale: str = DEFAULT_LOCALE
    active: bool = True


@dataclass
class ParseResult:
    phrases: list[ParsedPhrase] = field(default_factory=list)
    # (row number, reason) for rows that could not be used.
    skipped: list[tuple[int, str]] = field(default_factory=list)
    rows_read: int = 0


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_header(first_cell: str) -> bool:
    return first_cell.strip().lower() in _HEADER_CELLS


def _rows_from_xlsx(data: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover - dependency is declared
        raise PhraseImportError("Spreadsheet support is not installed on the server.") from e
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise PhraseImportError(f"Could not read this .xlsx file: {e}") from e
    ws = wb.active
    out = [[_cell(c) for c in row[:3]] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return out


def _rows_from_xls(data: bytes) -> list[list[str]]:
    try:
        import xlrd
    except ImportError as e:  # pragma: no cover - dependency is declared
        raise PhraseImportError("Legacy .xls support is not installed on the server.") from e
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as e:
        raise PhraseImportError(f"Could not read this .xls file: {e}") from e
    sheet = book.sheet_by_index(0)
    return [
        [_cell(sheet.cell_value(r, c)) for c in range(min(3, sheet.ncols))]
        for r in range(sheet.nrows)
    ]


def _rows_from_csv(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes anything
        raise PhraseImportError("Could not decode this CSV file.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[_cell(c) for c in row[:3]] for row in csv.reader(io.StringIO(text), dialect)]


def parse_phrase_file(filename: str, data: bytes) -> ParseResult:
    """Parse an uploaded spreadsheet into phrases plus per-row skip reasons."""
    if not data:
        raise PhraseImportError("The uploaded file is empty.")

    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        rows = _rows_from_xlsx(data)
    elif name.endswith(".xls"):
        # Some tools name an .xlsx file ".xls"; fall back rather than fail.
        try:
            rows = _rows_from_xls(data)
        except PhraseImportError:
            if data[:2] == b"PK":
                rows = _rows_from_xlsx(data)
            else:
                raise
    elif name.endswith(".csv") or name.endswith(".txt"):
        rows = _rows_from_csv(data)
    elif data[:2] == b"PK":  # zip container -> xlsx
        rows = _rows_from_xlsx(data)
    else:
        raise PhraseImportError(
            "Unsupported file type. Upload a .xlsx, .xls or .csv file."
        )

    result = ParseResult(rows_read=len(rows))
    seen: set[str] = set()

    for index, row in enumerate(rows, start=1):
        cells = list(row) + ["", "", ""]
        text = cells[0].strip()

        if index == 1 and _is_header(text):
            continue
        if not text:
            continue  # blank padding rows are not worth reporting

        if len(text) < MIN_LEN:
            result.skipped.append((index, f"too short (minimum {MIN_LEN} characters)"))
            continue
        if len(text) > MAX_LEN:
            result.skipped.append((index, f"too long (maximum {MAX_LEN} characters)"))
            continue

        key = " ".join(text.lower().split())
        if key in seen:
            result.skipped.append((index, "duplicate of an earlier row in this file"))
            continue
        seen.add(key)

        locale = cells[1].strip() or DEFAULT_LOCALE
        active_raw = cells[2].strip().lower()
        active = active_raw not in _FALSEY if active_raw else True

        result.phrases.append(ParsedPhrase(row=index, text=text, locale=locale, active=active))

    return result
